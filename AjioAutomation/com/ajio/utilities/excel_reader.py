import openpyxl

def get_data(file_path, sheet_name):
    data_list = []
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    sheet = workbook[sheet_name]
    header_row = [cell.value for cell in next(sheet.iter_rows(max_row=1))]

    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_data = {}
        for key, cell in zip(header_row, row):
            if cell is None:
                value = ""
            elif isinstance(cell, bool):
                value = str(cell)
            elif isinstance(cell, (int, float)):
                value = str(int(cell)) if isinstance(cell, float) and cell.is_integer() else str(cell)
            else:
                value = str(cell)
            row_data[key] = value
        data_list.append(row_data)

    return data_list